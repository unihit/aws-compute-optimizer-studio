"""
Multi-Format Report Generator for AWS Compute Optimizer Studio.
Generates:
1. Multi-Sheet Styled Excel (.xlsx) via openpyxl
2. Executive & Engineering Word Report (.docx) via python-docx
3. Printable Executive HTML/PDF Report (.html)
"""

import io
from datetime import datetime
from typing import Dict, Any, List
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT


def generate_excel_report(data: Dict[str, Any], krw_rate: float = 1380.0) -> bytes:
    wb = openpyxl.Workbook()

    # Colors & Styles
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    sub_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    highlight_fill = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")
    risk_fill = PatternFill(start_color="ECFEFF", end_color="ECFEFF", fill_type="solid")

    header_font = Font(name="Malgun Gothic", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Malgun Gothic", size=15, bold=True, color="0F172A")
    bold_font = Font(name="Malgun Gothic", size=10, bold=True, color="0F172A")
    reg_font = Font(name="Malgun Gothic", size=10, color="1E293B")
    savings_font = Font(name="Malgun Gothic", size=10, bold=True, color="047857")

    thin_side = Side(border_style="thin", color="CBD5E1")
    border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    meta = data.get("metadata", {})
    kpis = data.get("kpis", {})
    items = data.get("items", [])

    # Sheet 1: Executive Summary
    ws_sum = wb.active
    ws_sum.title = "Executive_Summary"
    ws_sum.views.sheetView[0].showGridLines = True

    fx_info = data.get("exchangeRateInfo", {})
    fx_source = fx_info.get("source", "AWS Invoice / 기준환율")

    ws_sum["A1"] = "AWS Compute Optimizer 비용 최적화 진단 리포트 (Executive Summary)"
    ws_sum["A1"].font = title_font
    ws_sum["A2"] = (
        f"계정 ID: {meta.get('accountId')} ({meta.get('accountAlias')}) | "
        f"리전: {meta.get('region')} | 진단 기준일: {meta.get('collectedAt', '')[:10]} | "
        f"적용 환율: 1 USD = {krw_rate:,.2f} KRW ({fx_source})"
    )
    ws_sum["A2"].font = reg_font

    # KPI Summary Table
    kpi_headers = ["핵심 지표 (KPI)", "값 (USD / 수량)", "원화 환산 (KRW)", "비고"]
    for col_idx, text in enumerate(kpi_headers, start=1):
        cell = ws_sum.cell(row=4, column=col_idx, value=text)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    monthly_usd = kpis.get("totalMonthlySavingsUsd", 0.0)
    annual_usd = kpis.get("totalAnnualSavingsUsd", 0.0)
    quick_usd = kpis.get("quickWinMonthlySavingsUsd", 0.0)
    co_status = meta.get("computeOptimizerStatus", "Active")
    co_note = meta.get("computeOptimizerNote", "")
    fx_info = data.get("exchangeRateInfo", {})
    fx_source = fx_info.get("source", "AWS Invoice / 기준환율")
    ce = data.get("costExplorer", {})
    last_bill = ce.get("lastMonthTotalUsd", 0.0)
    bill_pct = ce.get("savingsVsBillPercentage")

    kpi_rows = [
        ("Compute Optimizer 가입 상태", f"{co_status}", "-", co_note or "Active 상태 (정상 수집)"),
        ("최근 30일 총 청구액 (Cost Explorer)", f"${last_bill:,.2f}" if last_bill > 0 else "-", f"₩{last_bill * krw_rate:,.0f}" if last_bill > 0 else "-", f"총 청구액 대비 약 {bill_pct}% 절감 가능" if bill_pct else "Cost Explorer 연동 기준"),
        ("분석 대상 전체 리소스 수", f"{kpis.get('totalResourcesAnalyzed', 0)} 개", "-", "EC2, EBS, RDS, ECS, Lambda, Idle, Extended Support, 네트워크, 스토리지 통합"),
        ("비용 절감 권고 리소스 수", f"{kpis.get('optimizableResourcesCount', 0)} 개", "-", "Overprovisioned / NotOptimized / Idle / ExtendedSupport"),
        ("EKS / RDS Extended Support 대상", f"{kpis.get('extendedSupportCount', 0)} 건 (${kpis.get('extendedSupportMonthlyUsd', 0.0):,.2f}/월)", f"₩{kpis.get('extendedSupportMonthlyUsd', 0.0) * krw_rate:,.0f} / 월", "표준 지원 종료 버전 추가 과금 (업그레이드 시 100% 제거)"),
        ("EIP / ALB / NAT 네트워크 점검", f"미사용 EIP {kpis.get('unattachedEipCount', 0)}/{kpis.get('totalEipCount', 0)}개 | 유휴 ALB {kpis.get('idleOrNoTargetAlbCount', 0)}/{kpis.get('totalAlbCount', 0)}개 | NAT {kpis.get('idleOrRedundantNatGwCount', 0)}/{kpis.get('totalNatGwCount', 0)}개", "-", "미연결 공인 IPv4, 유휴 ALB, NAT GW 및 S3 Endpoint 점검"),
        ("노후 스냅샷 / 로그 / S3 수명주기", f"90일+ 스냅샷 {kpis.get('agedSnapshotCount', 0)}개 ({kpis.get('agedSnapshotTotalGb', 0)}GB) | 무기한 로그 {kpis.get('neverExpireLogGroupCount', 0)}개", f"₩{kpis.get('storageLifecycleMonthlyUsd', 0.0) * krw_rate:,.0f} / 월", "장기 방치 EBS/RDS 스냅샷, CW Logs 보관주기, S3 Lifecycle"),
        ("성능 병목 위험 리소스 수", f"{kpis.get('performanceRiskCount', 0)} 개", "-", "Underprovisioned (스케일업 검토 필요)"),
        ("예상 월간 총 절감액 (Monthly)", f"${monthly_usd:,.2f}", f"₩{monthly_usd * krw_rate:,.0f}", "Compute Optimizer + FinOps 전체 통합 기준"),
        ("예상 연간 총 절감액 (Annual)", f"${annual_usd:,.2f}", f"₩{annual_usd * krw_rate:,.0f}", "연간 12개월 환산 기대 효과"),
        ("즉시 실행 가능 절감액 (Quick Wins)", f"${quick_usd:,.2f} / 월", f"₩{quick_usd * krw_rate:,.0f} / 월", "마이그레이션 난이도 Low 이하 즉시 조치분"),
    ]

    for r_idx, row_vals in enumerate(kpi_rows, start=5):
        for c_idx, val in enumerate(row_vals, start=1):
            c = ws_sum.cell(row=r_idx, column=c_idx, value=val)
            c.font = savings_font if "절감액" in row_vals[0] and c_idx in (2, 3) else reg_font
            c.border = border
            if "절감액" in row_vals[0]:
                c.fill = highlight_fill

    # Category Breakdown Table
    ws_sum.cell(row=18, column=1, value="리소스 카테고리별 절감 기회 요약").font = bold_font
    cat_headers = ["리소스 유형", "분석 리소스 수", "최적화 대상 수", "월간 예상 절감액 (USD)", "연간 예상 절감액 (USD)", "연간 예상 절감액 (KRW)"]
    for col_idx, text in enumerate(cat_headers, start=1):
        c = ws_sum.cell(row=19, column=col_idx, value=text)
        c.fill = sub_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center")
        c.border = border

    row_cursor = 20
    for cat_key, cat_obj in data.get("byCategory", {}).items():
        m_usd = cat_obj.get("monthlySavingsUsd", 0.0)
        a_usd = cat_obj.get("annualSavingsUsd", 0.0)
        vals = [
            cat_key,
            cat_obj.get("count", 0),
            cat_obj.get("optimizableCount", 0),
            m_usd,
            a_usd,
            round(a_usd * krw_rate),
        ]
        for c_idx, val in enumerate(vals, start=1):
            c = ws_sum.cell(row=row_cursor, column=c_idx, value=val)
            c.font = reg_font
            c.border = border
            if c_idx in (4, 5):
                c.number_format = '"$"#,##0.00'
            elif c_idx == 6:
                c.number_format = '"₩"#,##0'
        row_cursor += 1

    for col in range(1, 7):
        ws_sum.column_dimensions[get_column_letter(col)].width = 26

    # Helper to populate detailed sheets
    def _create_detail_sheet(title: str, filtered_items: List[Dict[str, Any]]):
        ws = wb.create_sheet(title=title)
        headers = [
            "리소스 카테고리",
            "리소스명 (Name)",
            "리소스 ID / ARN",
            "환경 (Env)",
            "담당 조직 (Owner)",
            "진단 결과 (Finding)",
            "현재 스펙 (Current)",
            "권고 스펙 (Recommended)",
            "현재 사용률 요약",
            "적용 난이도",
            "월간 절감액 (USD)",
            "연간 절감액 (USD)",
            "연간 절감액 (KRW)",
            "절감률 (%)",
            "세부 실행 가이드 (Action Plan)",
        ]
        for col_idx, h in enumerate(headers, start=1):
            c = ws.cell(row=1, column=col_idx, value=h)
            c.fill = header_fill
            c.font = header_font
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = border

        for r_idx, item in enumerate(filtered_items, start=2):
            m_usd = item.get("monthlySavingsUsd", 0.0)
            a_usd = item.get("annualSavingsUsd", 0.0)
            f_norm = str(item.get("finding", "")).strip().upper().replace("_", "")
            is_opt_same = f_norm == "OPTIMIZED" or (
                item.get("currentConfig") == item.get("recommendedConfig") and m_usd <= 0 and f_norm != "UNDERPROVISIONED"
            )
            if is_opt_same:
                rec_str = "적정 유지 (변경 불필요)"
            else:
                rec_str = item.get("recommendedConfig") or ""
                if item.get("alternativeConfig"):
                    rec_str += f" (x86 대안: {item.get('alternativeConfig')})"
            guidance_str = item.get("actionGuidance") or ""
            if item.get("compatibilityWarning"):
                guidance_str += f" | {item.get('compatibilityWarning')}"

            row_data = [
                item.get("category"),
                item.get("name"),
                item.get("id"),
                item.get("environment"),
                item.get("owner"),
                item.get("finding"),
                item.get("currentConfig"),
                rec_str,
                item.get("utilizationSummary"),
                item.get("migrationEffort"),
                m_usd,
                a_usd,
                round(a_usd * krw_rate),
                item.get("savingsPercentage", 0.0),
                guidance_str,
            ]
            for c_idx, val in enumerate(row_data, start=1):
                c = ws.cell(row=r_idx, column=c_idx, value=val)
                c.font = reg_font
                c.border = border
                if r_idx % 2 == 1:
                    c.fill = zebra_fill
                if item.get("finding") == "Underprovisioned":
                    c.fill = risk_fill
                elif m_usd >= 200 and c_idx in (11, 12, 13):
                    c.fill = highlight_fill
                    c.font = savings_font

                if c_idx in (11, 12):
                    c.number_format = '"$"#,##0.00'
                elif c_idx == 13:
                    c.number_format = '"₩"#,##0'
                elif c_idx == 14:
                    c.number_format = '0.0"%"'

        widths = [14, 26, 28, 14, 20, 18, 28, 28, 30, 14, 18, 18, 20, 12, 52]
        for idx, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = w

    # Sheet 2: All Recommendations Ranked
    _create_detail_sheet("All_Recommendations", items)
    # Sheet 3: EC2 & EBS
    _create_detail_sheet("EC2_and_EBS", [i for i in items if i["category"] in ("EC2", "EBS")])
    # Sheet 4: RDS & Serverless (ECS, Lambda)
    _create_detail_sheet("RDS_ECS_Lambda", [i for i in items if i["category"] in ("RDS", "ECS", "Lambda")])
    # Sheet 5: Idle Resources
    _create_detail_sheet("Idle_Resources", [i for i in items if i["category"] == "Idle"])
    # Sheet 6: Extended Support, Network (EIP/ALB/NAT) & Storage/Log Lifecycle
    _create_detail_sheet(
        "Extended_Network_Storage",
        [i for i in items if i["category"] in ("ExtendedSupport", "Network", "StorageLifecycle")],
    )

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def generate_word_report(data: Dict[str, Any], krw_rate: float = 1380.0) -> bytes:
    """
    Generates a formal Korean Executive & Technical Assessment Word (.docx) Report.
    """
    doc = Document()

    # Base style setup
    style = doc.styles["Normal"]
    style.font.name = "Malgun Gothic"
    style.font.size = Pt(10)

    meta = data.get("metadata", {})
    kpis = data.get("kpis", {})
    items = data.get("items", [])
    action_plan = data.get("actionPlan", {})

    monthly_usd = kpis.get("totalMonthlySavingsUsd", 0.0)
    annual_usd = kpis.get("totalAnnualSavingsUsd", 0.0)
    quick_usd = kpis.get("quickWinMonthlySavingsUsd", 0.0)
    co_status = meta.get("computeOptimizerStatus", "Active")
    co_status_label = (
        f"활성화 완료 ({co_status})"
        if co_status.lower() == "active"
        else f"⚠️ 미가입 ({co_status}) — 콘솔 Opt-In 활성화 후 14일 뒤 정밀 권고 재수집 권장"
    )

    # Title Section
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_sub = title_p.add_run("AWS Well-Architected Cost Optimization Pillar 기반\n")
    run_sub.font.size = Pt(11)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(37, 99, 235)

    run_title = title_p.add_run("AWS Compute Optimizer 리소스 비용 최적화 진단 보고서")
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(15, 23, 42)

    fx_info = data.get("exchangeRateInfo", {})
    fx_source = fx_info.get("source", "AWS Invoice / 기준환율")

    graviton_mode_str = (
        "EC2 Graviton(ARM64) 전환 제안 포함 (호환성 주의 표시)"
        if meta.get("allowEc2Graviton", True)
        else "EC2 동일 아키텍처(x86_64) 리사이징 전용 (Graviton 제외)"
    )
    meta_p = doc.add_paragraph(
        f"• 진단 대상 계정: {meta.get('accountId')} ({meta.get('accountAlias')})\n"
        f"• 대상 리전: {meta.get('region')} (최근 {meta.get('lookBackPeriodInDays', 14)}일 CloudWatch 메트릭 기준)\n"
        f"• Compute Optimizer 가입 상태: {co_status_label}\n"
        f"• EC2 아키텍처 권고 기준: {graviton_mode_str}\n"
        f"• 보고서 생성일: {datetime.now().strftime('%Y-%m-%d %H:%M')}\n"
        f"• 기준 환율: 1 USD = {krw_rate:,.2f} KRW (출처: {fx_source})"
    )
    meta_p.paragraph_format.space_after = Pt(14)

    # 1. 경영진 요약 (Executive Summary)
    h1 = doc.add_heading("1. 경영진 요약 (Executive Summary)", level=1)
    for r in h1.runs:
        r.font.color.rgb = RGBColor(15, 23, 42)

    if co_status.lower() != "active":
        warn_p = doc.add_paragraph()
        warn_r = warn_p.add_run(
            f"[안내] 현재 대상 계정은 AWS Compute Optimizer가 미가입({co_status}) 상태입니다. "
            f"즉시 진단 가능한 EKS/RDS Extended Support, 미사용 EIP, 유휴 ALB/NAT GW, 노후 스냅샷 및 로그 보존주기 항목을 우선 점검하였으며, "
            f"Compute Optimizer를 활성화(무료 Opt-In)하신 뒤 14일 후 재수집하시면 머신러닝 기반 EC2/RDS/ECS/Lambda 라이트사이징 권고가 추가 반영됩니다."
        )
        warn_r.bold = True
        warn_r.font.color.rgb = RGBColor(180, 83, 9)

    doc.add_paragraph(
        f"본 보고서는 AWS Compute Optimizer의 머신러닝 기반 리소스 사용률 분석 및 FinOps 확장 스캔 결과를 바탕으로 "
        f"현재 운영 중인 클라우드 인프라(EC2, EBS, RDS/Aurora, ECS Fargate, Lambda, 유휴 리소스, Extended Support, 네트워크 및 스토리지 수명주기)의 "
        f"과다 프로비저닝(Over-provisioned) 비용 낭비 요소와 성능 병목 위험(Under-provisioned)을 종합 진단하였습니다.\n\n"
        f"진단 결과, 총 {kpis.get('totalResourcesAnalyzed', 0)}개의 분석 대상 리소스 중 "
        f"{kpis.get('optimizableResourcesCount', 0)}개 리소스에서 즉각적인 비용 절감 기회가 식별되었으며, "
        f"권고안 이행 시 월간 약 ${monthly_usd:,.2f} (약 ₩{monthly_usd * krw_rate:,.0f}), "
        f"연간 약 ${annual_usd:,.2f} (약 ₩{annual_usd * krw_rate:,.0f})의 직접적인 인프라 비용 절감이 가능합니다."
    )

    # KPI Summary Table
    kpi_table = doc.add_table(rows=1, cols=4)
    kpi_table.style = "Table Grid"
    kpi_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = kpi_table.rows[0].cells
    hdr_titles = ["구분 지표", "진단 수치 / 월간 금액", "연간 환산 효과 (KRW)", "핵심 시사점"]
    for i, t in enumerate(hdr_titles):
        hdr_cells[i].text = t
        for p in hdr_cells[i].paragraphs:
            for r in p.runs:
                r.font.bold = True

    fx_info = data.get("exchangeRateInfo", {})
    fx_source = fx_info.get("source", "AWS Invoice / 기준환율")
    ce = data.get("costExplorer", {})
    bill_usd = ce.get("lastMonthTotalUsd", 0.0)
    bill_pct = ce.get("savingsVsBillPercentage")
    sp_cov = ce.get("savingsPlansCoveragePct")
    bill_str = f"${bill_usd:,.2f} (청구액 대비 {bill_pct}% 절감 가능)" if bill_usd > 0 else "Cost Explorer 미연동"

    summary_rows = [
        ("총 예상 절감 가능액", f"${monthly_usd:,.2f} / 월", f"₩{annual_usd * krw_rate:,.0f} / 년", "전체 권고안(Rank 1) 및 버전/네트워크/스토리지 합산"),
        ("최근 30일 AWS 총 청구액", bill_str, f"SP 커버리지 {sp_cov}%" if sp_cov is not None else "-", "Cost Explorer (UnblendedCost) 기준"),
        ("Quick Win (즉시 조치 가능)", f"${quick_usd:,.2f} / 월", f"₩{quick_usd * 12 * krw_rate:,.0f} / 년", "아키텍처 변경 없이 즉시 리사이징/삭제 가능"),
        ("EKS / RDS Extended Support", f"{kpis.get('extendedSupportCount', 0)} 건 (${kpis.get('extendedSupportMonthlyUsd', 0.0):,.2f}/월)", f"₩{kpis.get('extendedSupportMonthlyUsd', 0.0) * 12 * krw_rate:,.0f} / 년", "표준 지원 종료 버전 추가 요금 (업그레이드 권고)"),
        ("EIP / ALB / NAT 네트워크 점검", f"미사용 EIP {kpis.get('unattachedEipCount', 0)}/{kpis.get('totalEipCount', 0)}개 | 유휴 ALB {kpis.get('idleOrNoTargetAlbCount', 0)}/{kpis.get('totalAlbCount', 0)}개 | NAT {kpis.get('idleOrRedundantNatGwCount', 0)}/{kpis.get('totalNatGwCount', 0)}개", "-", "미연결 IPv4 반환, 유휴 ALB 및 S3 Gateway 엔드포인트 점검"),
        ("노후 스냅샷 / CW 로그 / S3", f"90일+ 스냅샷 {kpis.get('agedSnapshotCount', 0)}개 ({kpis.get('agedSnapshotTotalGb', 0)}GB) | 로그 영구보존 {kpis.get('neverExpireLogGroupCount', 0)}개 | S3 {kpis.get('s3NoLifecycleBucketCount', 0)}개", f"₩{kpis.get('storageLifecycleMonthlyUsd', 0.0) * 12 * krw_rate:,.0f} / 년", "장기 방치 백업 정리 및 보존주기(Retention/Lifecycle) 설정"),
        ("비용 최적화 대상 리소스", f"{kpis.get('optimizableResourcesCount', 0)} 개 / 총 {kpis.get('totalResourcesAnalyzed', 0)} 개", "-", "과다 사양(Overprovisioned) 및 유휴(Idle)"),
        ("성능 위험(Underprovisioned)", f"{kpis.get('performanceRiskCount', 0)} 개", "안정성 확보 필수", "CPU/메모리 90% 이상 피크 발생 리소스"),
    ]
    for r_data in summary_rows:
        row_cells = kpi_table.add_row().cells
        for i, val in enumerate(r_data):
            row_cells[i].text = str(val)

    doc.add_paragraph()

    # 2. 서비스별 비용 절감 기회 분석
    h2 = doc.add_heading("2. AWS 리소스 유형별 진단 현황", level=1)
    for r in h2.runs:
        r.font.color.rgb = RGBColor(15, 23, 42)

    cat_table = doc.add_table(rows=1, cols=5)
    cat_table.style = "Table Grid"
    cat_hdr = ["리소스 카테고리", "점검 수량", "최적화 대상", "월간 절감액 (USD)", "연간 절감액 (KRW)"]
    for i, t in enumerate(cat_hdr):
        cat_table.rows[0].cells[i].text = t
        for p in cat_table.rows[0].cells[i].paragraphs:
            for r in p.runs:
                r.font.bold = True

    for cat_key, cat_obj in data.get("byCategory", {}).items():
        row_cells = cat_table.add_row().cells
        m_val = cat_obj.get("monthlySavingsUsd", 0.0)
        a_val = cat_obj.get("annualSavingsUsd", 0.0)
        row_cells[0].text = cat_key
        row_cells[1].text = f"{cat_obj.get('count', 0)} 개"
        row_cells[2].text = f"{cat_obj.get('optimizableCount', 0)} 개"
        row_cells[3].text = f"${m_val:,.2f}"
        row_cells[4].text = f"₩{a_val * krw_rate:,.0f}"

    doc.add_paragraph()

    # 3. 단계별 실행 로드맵 (3-Phase Action Plan)
    h3 = doc.add_heading("3. 단계별 비용 최적화 실행 계획 (Action Plan)", level=1)
    for r in h3.runs:
        r.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_heading("Phase 1. 즉시 실행 과제 (Quick Wins - 무중단/저위험 최적화)", level=2)
    doc.add_paragraph(
        "유휴(Idle/Unattached) 리소스 정리, EBS gp2 → gp3 온라인 볼륨 전환, 동일 아키텍처(x86) 내 인스턴스 다운사이징 및 "
        "ECS Fargate / Lambda 메모리 설정 조정으로 서비스 영향도를 최소화하며 즉각적인 비용 절감을 달성합니다."
    )
    for item in action_plan.get("phase1QuickWins", [])[:6]:
        p = doc.add_paragraph(style="List Bullet")
        r_bold = p.add_run(f"[{item['category']}] {item['name']} ({item['id']}): ")
        r_bold.bold = True
        p.add_run(
            f"{item['currentConfig']} → {item['recommendedConfig']} | "
            f"월 ${item['monthlySavingsUsd']:,.2f} (연 ₩{item['annualSavingsUsd'] * krw_rate:,.0f}) 절감 — {item['actionGuidance']}"
        )

    doc.add_heading("Phase 2. 아키텍처 현대화 및 Graviton 전환 (Medium Effort)", level=2)
    doc.add_paragraph(
        "AWS Graviton(ARM64) 기반 인스턴스(m7g, r7g, db.r6g 등)로의 전환을 통해 가격 대비 성능을 극대화합니다. "
        "애플리케이션 런타임 호환성 검증 후 배포 파이프라인에 반영하는 것을 권장합니다."
    )
    arch_items = action_plan.get("phase2ArchModernization", [])
    if arch_items:
        for item in arch_items:
            p = doc.add_paragraph(style="List Bullet")
            r_bold = p.add_run(f"[{item['category']}] {item['name']} ({item['id']}): ")
            r_bold.bold = True
            p.add_run(
                f"{item['currentConfig']} → {item['recommendedConfig']} | "
                f"월 ${item['monthlySavingsUsd']:,.2f} 절감 ({item['utilizationSummary']})"
            )
    else:
        doc.add_paragraph("• 현재 별도의 아키텍처(ARM64 등) 마이그레이션이 필요한 고난이도 항목이 없습니다.")

    doc.add_heading("Phase 3. 성능 병목 해소 및 안정성 강화 (Under-provisioned 리소스)", level=2)
    risk_items = action_plan.get("phase3PerformanceStabilization", [])
    if risk_items:
        for item in risk_items:
            p = doc.add_paragraph(style="List Bullet")
            r_bold = p.add_run(f"[주의 - {item['category']}] {item['name']} ({item['id']}): ")
            r_bold.bold = True
            r_bold.font.color.rgb = RGBColor(185, 28, 28)
            p.add_run(f"{item['utilizationSummary']} → {item['actionGuidance']}")
    else:
        doc.add_paragraph("• 성능 병목(Under-provisioned)으로 분류된 위험 리소스가 없습니다.")

    # 4. 전체 리소스 상세 권고 명세표
    h4 = doc.add_heading("4. 리소스별 상세 최적화 권고 명세", level=1)
    for r in h4.runs:
        r.font.color.rgb = RGBColor(15, 23, 42)

    det_table = doc.add_table(rows=1, cols=6)
    det_table.style = "Table Grid"
    det_hdr = ["유형", "리소스명 / ID", "진단 상태", "현재 스펙 → 권고 스펙", "사용률 지표", "월간 절감액(USD)"]
    for i, t in enumerate(det_hdr):
        det_table.rows[0].cells[i].text = t
        for p in det_table.rows[0].cells[i].paragraphs:
            for r in p.runs:
                r.font.bold = True

    for item in items:
        rc = det_table.add_row().cells
        rc[0].text = item["category"]
        rc[1].text = f"{item['name']}\n({item['id']})"
        rc[2].text = item["finding"]
        f_norm = str(item.get("finding", "")).strip().upper().replace("_", "")
        is_opt_same = f_norm == "OPTIMIZED" or (
            item.get("currentConfig") == item.get("recommendedConfig")
            and item.get("monthlySavingsUsd", 0.0) <= 0
            and f_norm != "UNDERPROVISIONED"
        )
        if is_opt_same:
            spec_txt = f"현재: {item['currentConfig']}\n→ [적정 유지] 스펙 변경 불필요"
        else:
            spec_txt = f"{item['currentConfig']}\n→ {item['recommendedConfig']}"
            if item.get("alternativeConfig"):
                spec_txt += f"\n(x86 대안: {item['alternativeConfig']})"
            if item.get("compatibilityWarning"):
                spec_txt += f"\n{item['compatibilityWarning']}"
        rc[3].text = spec_txt
        rc[4].text = item["utilizationSummary"]
        rc[5].text = (
            "-"
            if item.get("monthlySavingsUsd", 0.0) <= 0
            else f"${item['monthlySavingsUsd']:,.2f}\n({item['savingsPercentage']}%)"
        )

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


def generate_executive_html_report(data: Dict[str, Any], krw_rate: float = 1380.0) -> str:
    """
    Generates a standalone, print-ready Executive HTML/PDF Report with clean styling.
    """
    meta = data.get("metadata", {})
    kpis = data.get("kpis", {})
    items = data.get("items", [])

    monthly_usd = kpis.get("totalMonthlySavingsUsd", 0.0)
    annual_usd = kpis.get("totalAnnualSavingsUsd", 0.0)
    quick_usd = kpis.get("quickWinMonthlySavingsUsd", 0.0)
    co_status = meta.get("computeOptimizerStatus", "Active")
    co_banner_html = ""
    if co_status.lower() != "active":
        co_banner_html = f"""
        <div style="margin-bottom:20px; padding:14px 18px; border-radius:8px; background:#fffbeb; border:1px solid #f59e0b; color:#92400e; font-size:13px;">
            <strong>⚠️ AWS Compute Optimizer 미가입 ({co_status}) 안내:</strong>
            현재 계정은 Compute Optimizer가 비활성화 상태입니다. 즉시 진단 가능한 Extended Support, 미사용 EIP, 유휴 ALB/NAT GW, 노후 스냅샷 및 로그 보존주기는 정상 수집되었으며,
            AWS 콘솔에서 Compute Optimizer를 <strong>무료 Opt-In(활성화)</strong>하신 뒤 <strong>14일 후</strong> 재수집하시면 EC2/EBS/RDS/ECS/Lambda 머신러닝 라이트사이징 권고가 함께 포함됩니다.
        </div>
        """
    fx_info = data.get("exchangeRateInfo", {})
    fx_source = fx_info.get("source", "AWS Invoice / 기준환율")
    ce = data.get("costExplorer", {})
    bill_usd = ce.get("lastMonthTotalUsd", 0.0)
    bill_pct = ce.get("savingsVsBillPercentage")
    bill_sub_html = (
        f'<div style="font-size:11px; color:#ea580c; font-weight:700; margin-top:4px;">최근 30일 청구액(${bill_usd:,.0f}) 대비 -{bill_pct}% 절감 가능</div>'
        if bill_usd > 0 and bill_pct is not None
        else ""
    )

    rows_html = ""
    for item in items:
        f_norm = str(item.get("finding", "")).strip().upper().replace("_", "")
        is_opt_same = f_norm == "OPTIMIZED" or (
            item.get("currentConfig") == item.get("recommendedConfig")
            and item.get("monthlySavingsUsd", 0.0) <= 0
            and f_norm != "UNDERPROVISIONED"
        )
        badge_color = (
            "#ea580c" if f_norm == "UNDERPROVISIONED"
            else "#0891b2" if f_norm == "OVERPROVISIONED"
            else "#ca8a04" if f_norm == "NOTOPTIMIZED"
            else "#d97706" if item["category"] == "Idle" or f_norm in ("IDLE", "UNATTACHED")
            else "#059669" if is_opt_same
            else "#7c3aed" if f_norm == "EXTENDEDSUPPORT"
            else "#0891b2"
        )
        badge_label = (
            "🔥 Underprovisioned" if f_norm == "UNDERPROVISIONED"
            else "🧊 Overprovisioned" if f_norm == "OVERPROVISIONED"
            else "⚠️ NotOptimized" if f_norm == "NOTOPTIMIZED"
            else "✅ Optimized (적정)" if is_opt_same
            else "⏳ Extended Support" if f_norm == "EXTENDEDSUPPORT"
            else f"🧹 {item['finding']} (유휴)" if (item["category"] == "Idle" or f_norm in ("IDLE", "UNATTACHED"))
            else item["finding"]
        )
        spec_html = (
            f'<div style="margin-top:3px;"><span style="display:inline-block; background:#d1fae5; color:#065f46; border:1px solid #6ee7b7; font-size:10px; font-weight:700; padding:1px 6px; border-radius:4px; margin-right:4px;">적정 유지</span><span style="color:#059669; font-weight:700; font-size:12px;">스펙 변경 불필요 (현재 구성 최적)</span></div>'
            if is_opt_same
            else f'<div style="margin-top:3px;"><span style="display:inline-block; background:#dc2626; color:#fff; font-size:10px; font-weight:700; padding:1px 6px; border-radius:4px; margin-right:4px;">권고</span><span style="color:#ea580c; font-weight:800; font-size:12px;">{item["recommendedConfig"]}</span></div>'
        )
        sav_html = (
            '<span style="color:#94a3b8;">-</span>'
            if item.get("monthlySavingsUsd", 0.0) <= 0
            else f"""
                <div>
                    <span style="font-weight:800; font-size:14px; color:#ea580c;">${item['monthlySavingsUsd']:,.2f}</span><span style="font-size:11px; color:#64748b;">/월</span>
                    <span style="display:inline-block; margin-left:4px; padding:1px 7px; border-radius:99px; font-size:11px; font-weight:800; background:#ffedd5; color:#c2410c; border:1px solid #fdba74;">-{item['savingsPercentage']}%</span>
                </div>
                <div style="font-size:11px; color:#475569; font-weight:600; margin-top:3px;">₩{item['monthlySavingsUsd']*krw_rate:,.0f}/월 · 연 ₩{item['annualSavingsUsd']*krw_rate:,.0f}</div>
            """
        )
        rows_html += f"""
        <tr>
            <td><strong>{item['category']}</strong></td>
            <td>
                <div style="font-weight:600; color:#0f172a;">{item['name']}</div>
                <div style="font-size:11px; color:#64748b; font-family:monospace;">{item['id']} | {item['environment']} | {item['owner']}</div>
            </td>
            <td style="white-space:nowrap;"><span style="display:inline-block; white-space:nowrap; padding:3px 10px; border-radius:99px; font-size:11px; font-weight:700; color:#fff; background:{badge_color};">{badge_label}</span></td>
            <td>
                <div style="color:#475569; font-size:12px;">현재: {item['currentConfig']}</div>
                {spec_html}
                {f'<div style="margin-top:4px; padding:3px 6px; border-radius:4px; background:#fef3c7; border:1px solid #fcd34d; color:#92400e; font-size:10px; font-weight:600;">{item["compatibilityWarning"]}</div>' if item.get("compatibilityWarning") else ''}
            </td>
            <td style="font-size:12px; color:#334155;">{item['utilizationSummary']}</td>
            <td style="text-align:right; white-space:nowrap;">
                {sav_html}
            </td>
        </tr>
        """

    cat_cards = ""
    for cat_key, cat_obj in data.get("byCategory", {}).items():
        cat_cards += f"""
        <div style="border:1px solid #e2e8f0; border-radius:8px; padding:12px; background:#f8fafc;">
            <div style="font-size:12px; color:#64748b; font-weight:600;">{cat_key} ({cat_obj['optimizableCount']}/{cat_obj['count']} 대상)</div>
            <div style="font-size:18px; font-weight:700; color:#0f172a; margin-top:4px;">${cat_obj['monthlySavingsUsd']:,.2f}<span style="font-size:12px; font-weight:400; color:#64748b;"> /월</span></div>
            <div style="font-size:11px; color:#059669;">연간 ₩{cat_obj['annualSavingsUsd']*krw_rate:,.0f} 절감</div>
        </div>
        """

    return f"""<!DOCTYPE html>
<html lang="ko">
<head>
<meta charset="UTF-8">
<title>AWS Compute Optimizer 비용 최적화 리포트 - {meta.get('accountId')}</title>
<style>
    @page {{ size: A4 landscape; margin: 14mm; }}
    body {{ font-family: 'Pretendard', 'Malgun Gothic', -apple-system, sans-serif; color: #0f172a; margin: 0; padding: 32px; background: #ffffff; }}
    .no-print {{ margin-bottom: 20px; display: flex; justify-content: space-between; align-items: center; background: #0f172a; color: #fff; padding: 12px 20px; border-radius: 8px; }}
    .btn-print {{ background: #2563eb; color: #fff; border: none; padding: 8px 16px; border-radius: 6px; font-weight: 600; cursor: pointer; }}
    table {{ width: 100%; border-collapse: collapse; margin-top: 16px; }}
    th {{ background: #0f172a; color: #ffffff; text-align: left; padding: 10px 12px; font-size: 12px; }}
    td {{ border-bottom: 1px solid #e2e8f0; padding: 10px 12px; font-size: 13px; vertical-align: top; }}
    @media print {{ .no-print {{ display: none !important; }} body {{ padding: 0; }} }}
</style>
</head>
<body>
    <div class="no-print">
        <div><strong>AWS Compute Optimizer 공식 보고서 미리보기</strong> — 인쇄(Ctrl+P) 버튼을 눌러 PDF로 저장할 수 있습니다.</div>
        <button class="btn-print" onclick="window.print()">🖨️ PDF 저장 / 인쇄하기</button>
    </div>

    <div style="border-bottom: 3px solid #0f172a; padding-bottom: 16px; margin-bottom: 24px; display:flex; justify-content:space-between; align-items:flex-end;">
        <div>
            <div style="font-size:12px; font-weight:700; color:#2563eb; letter-spacing:0.05em;">AWS WELL-ARCHITECTED COST OPTIMIZATION REPORT</div>
            <h1 style="margin: 6px 0 0 0; font-size: 26px;">AWS Compute Optimizer 비용 최적화 진단 리포트</h1>
        </div>
        <div style="text-align:right; font-size:12px; color:#475569;">
            <div><strong>Account:</strong> {meta.get('accountId')} ({meta.get('accountAlias')}) | <strong>Compute Optimizer:</strong> {co_status}</div>
            <div><strong>Region:</strong> {meta.get('region')} | <strong>기준 환율:</strong> ₩{krw_rate:,.2f}/USD ({fx_source})</div>
            <div><strong>진단 일시:</strong> {meta.get('collectedAt', '')[:19].replace('T', ' ')} UTC</div>
        </div>
    </div>

    {co_banner_html}

    <div style="display:grid; grid-template-columns: repeat(4, 1fr); gap: 16px; margin-bottom: 24px;">
        <div style="border: 1px solid #cbd5e1; border-left: 5px solid #ea580c; border-radius: 8px; padding: 16px;">
            <div style="font-size:12px; color:#475569;">예상 월간 총 절감액</div>
            <div style="font-size:26px; font-weight:800; color:#ea580c; margin-top:4px;">${monthly_usd:,.2f}</div>
            <div style="font-size:13px; font-weight:600; color:#0f172a;">월 약 ₩{monthly_usd*krw_rate:,.0f}</div>
            {bill_sub_html}
        </div>
        <div style="border: 1px solid #cbd5e1; border-left: 5px solid #2563eb; border-radius: 8px; padding: 16px;">
            <div style="font-size:12px; color:#475569;">예상 연간 총 절감액 (12개월)</div>
            <div style="font-size:26px; font-weight:800; color:#1d4ed8; margin-top:4px;">${annual_usd:,.2f}</div>
            <div style="font-size:13px; font-weight:600; color:#0f172a;">연 약 ₩{annual_usd*krw_rate:,.0f}</div>
        </div>
        <div style="border: 1px solid #cbd5e1; border-left: 5px solid #0ea5e9; border-radius: 8px; padding: 16px;">
            <div style="font-size:12px; color:#475569;">Quick Win (즉시 조치 가능 절감액)</div>
            <div style="font-size:26px; font-weight:800; color:#0369a1; margin-top:4px;">${quick_usd:,.2f}<span style="font-size:14px;">/월</span></div>
            <div style="font-size:13px; color:#475569;">난이도 Low 이하 리소스</div>
        </div>
        <div style="border: 1px solid #cbd5e1; border-left: 5px solid #dc2626; border-radius: 8px; padding: 16px;">
            <div style="font-size:12px; color:#475569;">최적화 대상 / 성능 위험 리소스</div>
            <div style="font-size:26px; font-weight:800; color:#0f172a; margin-top:4px;">{kpis.get('optimizableResourcesCount', 0)}건 <span style="font-size:16px; color:#dc2626;">(위험 {kpis.get('performanceRiskCount', 0)}건)</span></div>
            <div style="font-size:13px; color:#475569;">전체 {kpis.get('totalResourcesAnalyzed', 0)}개 리소스 분석 완료</div>
        </div>
    </div>

    <h3 style="margin: 20px 0 10px 0; font-size: 15px;">서비스 카테고리별 절감 요약 (Extended Support · EIP/ALB/NAT · 노후 스냅샷/로그/S3 포함)</h3>
    <div style="display:grid; grid-template-columns: repeat(3, 1fr); gap: 12px; margin-bottom: 24px;">
        {cat_cards}
    </div>

    <h3 style="margin: 20px 0 8px 0; font-size: 15px;">리소스별 상세 최적화 권고 내역 (절감액 순)</h3>
    <table>
        <thead>
            <tr>
                <th style="width:70px;">유형</th>
                <th style="width:230px;">리소스명 / 식별자</th>
                <th style="width:110px;">진단 상태</th>
                <th style="width:250px;">스펙 변경 권고 (Current → Recommended)</th>
                <th>최근 14일 피크 사용률 및 진단 근거</th>
                <th style="width:140px; text-align:right;">예상 월 절감액</th>
            </tr>
        </thead>
        <tbody>
            {rows_html}
        </tbody>
    </table>
</body>
</html>
"""
